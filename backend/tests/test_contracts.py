from __future__ import annotations

from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from app.db import SessionLocal, engine, init_business_db
from app.main import app
from app.models import Base, IdempotencyRecord, ReviewRun, TaskJob


@pytest.fixture(autouse=True)
def clean_database() -> None:
    init_business_db()
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)


@pytest.fixture()
def client() -> TestClient:
    return TestClient(app)


def headers(user: str = "demo-rm", key: str = "key-001") -> dict[str, str]:
    return {"X-Demo-User-Id": user, "Idempotency-Key": key}


def create_case(client: TestClient) -> dict:
    response = client.post(
        "/api/v1/cases",
        headers=headers(),
        json={
            "case_no": "CASE-001",
            "customer_name": "合成科技有限公司",
            "customer_key": "SYNTH-001",
            "review_date": "2026-08-10",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _valid_material_bytes(suffix: str, term_months: int = 24) -> bytes:
    text = "\n".join(
        [
            "企业名称：合成科技有限公司",
            "统一社会信用代码：91310000SYNTH001",
            "成立日期：2020-01-01",
            "法定代表人：张三",
            "所属行业：制造业",
            "申请金额：1000000",
            "币种：CNY",
            f"申请期限：{term_months}个月",
            "贷款用途：流动资金",
            "总资产：5000000",
            "总负债：2000000",
            "流动资产：3000000",
            "流动负债：1500000",
            "营业收入：8000000",
            "净利润：600000",
        ]
    )
    if suffix == "pdf":
        import fitz

        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), text, fontsize=9)
        return document.tobytes()
    if suffix == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph(text)
        output = BytesIO()
        document.save(output)
        return output.getvalue()
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    if sheet is None:
        raise AssertionError("A workbook must have an active worksheet")
    for line in text.splitlines():
        key, value = line.split("：", 1)
        sheet.append([key, value])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def upload_required_documents(
    client: TestClient, case_id: str, valid: bool = False, conflict: bool = False
) -> list[str]:
    document_ids: list[str] = []
    types = ["BUSINESS_LICENSE", "CREDIT_APPLICATION", "DUE_DILIGENCE", "FINANCIAL_STATEMENTS"]
    suffixes = ["docx", "docx", "docx", "xlsx"]
    for index, (document_type, suffix) in enumerate(zip(types, suffixes, strict=True)):
        response = client.post(
            f"/api/v1/cases/{case_id}/documents",
            headers=headers(key=f"document-{index}"),
            data={"document_type": document_type},
            files={
                "file": (
                    f"material-{index}.{suffix}",
                    _valid_material_bytes(
                        suffix,
                        term_months=48 if conflict and document_type == "DUE_DILIGENCE" else 24,
                    )
                    if valid
                    else b"synthetic material",
                    "application/octet-stream",
                )
            },
        )
        assert response.status_code == 201, response.text
        document_ids.append(response.json()["id"])
    return document_ids


def test_health_and_openapi_contract(client: TestClient) -> None:
    assert client.get("/health").json()["status"] == "ok"
    openapi = client.get("/openapi.json").json()
    assert "/api/v1/cases/{case_id}/runs" in openapi["paths"]
    assert "/api/v1/runs/{run_id}/report/export" in openapi["paths"]


def test_case_document_run_contract_and_idempotency(client: TestClient) -> None:
    case = create_case(client)
    duplicate = client.post(
        "/api/v1/cases",
        headers=headers(key="key-001"),
        json={
            "case_no": "CASE-001",
            "customer_name": "合成科技有限公司",
            "customer_key": "SYNTH-001",
            "review_date": "2026-08-10",
        },
    )
    assert duplicate.status_code == 201
    assert duplicate.json()["id"] == case["id"]

    document_ids = upload_required_documents(client, case["id"], valid=True)
    detail = client.get(f"/api/v1/cases/{case['id']}", headers={"X-Demo-User-Id": "demo-rm"})
    assert detail.status_code == 200
    assert len(detail.json()["documents"]) == 4
    assert detail.json()["version"] == 5

    payload = {"document_version_ids": document_ids, "expected_case_version": 5}
    run = client.post(
        f"/api/v1/cases/{case['id']}/runs",
        headers=headers(key="run-001"),
        json=payload,
    )
    assert run.status_code == 202, run.text
    duplicate_run = client.post(
        f"/api/v1/cases/{case['id']}/runs",
        headers=headers(key="run-001"),
        json=payload,
    )
    assert duplicate_run.status_code == 202
    assert duplicate_run.json()["id"] == run.json()["id"]
    assert run.json()["status"] == "QUEUED"
    assert run.json()["progress_percent"] == 0

    with SessionLocal() as db:
        assert db.query(ReviewRun).count() == 1
        assert db.query(TaskJob).count() == 1
        assert db.query(IdempotencyRecord).count() == 6


def test_role_boundary_and_problem_details(client: TestClient) -> None:
    response = client.post(
        "/api/v1/cases",
        headers=headers(user="demo-reviewer", key="reviewer-case"),
        json={
            "case_no": "CASE-002",
            "customer_name": "合成企业",
            "customer_key": "SYNTH-002",
            "review_date": "2026-08-10",
        },
    )
    assert response.status_code == 403
    body = response.json()
    assert body["code"] == "FORBIDDEN"
    assert "trace_id" in body


def test_worker_claim_has_lease_and_state_contract() -> None:
    client = TestClient(app)
    case = create_case(client)
    document_ids = upload_required_documents(client, case["id"])
    response = client.post(
        f"/api/v1/cases/{case['id']}/runs",
        headers=headers(key="run-worker"),
        json={"document_version_ids": document_ids, "expected_case_version": 5},
    )
    assert response.status_code == 202
    from app.service import claim_job

    with SessionLocal() as db:
        claimed = claim_job(db, "test-worker")
        assert claimed is not None
        task, run = claimed
        assert task.status.value == "LEASED"
        assert task.leased_until is not None
        assert run.status.value == "RUNNING"


def test_worker_once_advances_to_report_gate() -> None:
    client = TestClient(app)
    case = create_case(client)
    document_ids = upload_required_documents(client, case["id"], valid=True)
    response = client.post(
        f"/api/v1/cases/{case['id']}/runs",
        headers=headers(key="run-worker-gate"),
        json={"document_version_ids": document_ids, "expected_case_version": 5},
    )
    assert response.status_code == 202

    from app.worker import Worker

    assert Worker(worker_id="test-worker", allow_memory_checkpoint=True).run_once() is True
    run_status = client.get(
        f"/api/v1/runs/{response.json()['id']}",
        headers={"X-Demo-User-Id": "demo-rm"},
    )
    assert run_status.status_code == 200
    assert run_status.json()["status"] == "WAITING_REPORT_REVIEW"
    results = client.get(
        f"/api/v1/runs/{response.json()['id']}/review-results",
        headers={"X-Demo-User-Id": "demo-reviewer"},
    )
    assert results.status_code == 200, results.text
    assert results.json()["summary_outcome"] == "PASS"
    assert len(results.json()["rules"]) == 10
    assert all(item["status"] == "PASS" for item in results.json()["rules"])


def test_material_fact_gate_human_selection_and_idempotency(client: TestClient) -> None:
    case = create_case(client)
    document_ids = upload_required_documents(client, case["id"], valid=True, conflict=True)
    response = client.post(
        f"/api/v1/cases/{case['id']}/runs",
        headers=headers(key="run-fact-gate"),
        json={"document_version_ids": document_ids, "expected_case_version": 5},
    )
    assert response.status_code == 202
    run_id = response.json()["id"]

    from app.worker import Worker

    assert Worker(worker_id="fact-worker", allow_memory_checkpoint=True).run_once() is True
    run_status = client.get(f"/api/v1/runs/{run_id}", headers={"X-Demo-User-Id": "demo-reviewer"})
    assert run_status.json()["status"] == "WAITING_FACT_REVIEW"

    facts_response = client.get(
        f"/api/v1/runs/{run_id}/facts", headers={"X-Demo-User-Id": "demo-reviewer"}
    )
    assert facts_response.status_code == 200, facts_response.text
    facts = facts_response.json()
    conflict = next(item for item in facts["conflicts"] if item["field"] == "F08")
    selected = next(item for item in conflict["candidates"] if item["normalized_value"] == 48)
    review_payload = {
        "expected_snapshot_version": facts["snapshot_version"],
        "decisions": [
            {
                "conflict_id": conflict["conflict_id"],
                "action": "SELECT_SOURCE",
                "selected_evidence_id": selected["evidence_id"],
            }
        ],
    }
    review_headers = headers(user="demo-reviewer", key="fact-review-001")
    reviewed = client.post(
        f"/api/v1/runs/{run_id}/fact-review",
        headers=review_headers,
        json=review_payload,
    )
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["status"] == "WAITING_REPORT_REVIEW"

    repeated = client.post(
        f"/api/v1/runs/{run_id}/fact-review",
        headers=review_headers,
        json=review_payload,
    )
    assert repeated.status_code == 200
    assert repeated.json()["id"] == run_id

    final_facts = client.get(
        f"/api/v1/runs/{run_id}/facts", headers={"X-Demo-User-Id": "demo-reviewer"}
    ).json()
    assert final_facts["requires_review"] is False
    assert final_facts["fields"]["F08"]["selected_value"] == 48
    results = client.get(
        f"/api/v1/runs/{run_id}/review-results",
        headers={"X-Demo-User-Id": "demo-reviewer"},
    )
    assert results.status_code == 200, results.text
    assert results.json()["summary_outcome"] == "NON_COMPLIANT"
    r07 = next(item for item in results.json()["rules"] if item["rule_id"] == "R07")
    assert r07["status"] == "FAIL"
    assert any(
        hit["chunk_id"] in r07["evidence_refs"]
        for hit in results.json()["retrieval"]["by_rule"]["R07"]
        if hit["selected"]
    )


def test_report_review_confirm_visibility_and_markdown_export(client: TestClient) -> None:
    case = create_case(client)
    document_ids = upload_required_documents(client, case["id"], valid=True)
    response = client.post(
        f"/api/v1/cases/{case['id']}/runs",
        headers=headers(key="run-report-confirm"),
        json={"document_version_ids": document_ids, "expected_case_version": 5},
    )
    assert response.status_code == 202
    run_id = response.json()["id"]

    from app.worker import Worker

    assert Worker(worker_id="report-worker", allow_memory_checkpoint=True).run_once() is True
    reviewer_headers = {"X-Demo-User-Id": "demo-reviewer"}
    draft = client.get(f"/api/v1/runs/{run_id}/report", headers=reviewer_headers)
    assert draft.status_code == 200, draft.text
    draft_payload = draft.json()
    assert draft_payload["report_status"] == "AWAITING_REVIEW"
    assert "授信智能合规审查报告" in draft_payload["markdown"]
    assert "不构成授信审批" in draft_payload["markdown"]

    rm_draft = client.get(f"/api/v1/runs/{run_id}/report", headers={"X-Demo-User-Id": "demo-rm"})
    assert rm_draft.status_code == 403
    assert rm_draft.json()["code"] == "REPORT_NOT_CONFIRMED"

    review_payload = {
        "expected_snapshot_version": draft_payload["snapshot_version"],
        "action": "CONFIRM_DRAFT",
    }
    review_headers = headers(user="demo-reviewer", key="report-review-001")
    confirmed = client.post(
        f"/api/v1/runs/{run_id}/report-review",
        headers=review_headers,
        json=review_payload,
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["status"] == "COMPLETED"

    repeated = client.post(
        f"/api/v1/runs/{run_id}/report-review",
        headers=review_headers,
        json=review_payload,
    )
    assert repeated.status_code == 200
    assert repeated.json()["id"] == run_id

    final_report = client.get(
        f"/api/v1/runs/{run_id}/report", headers={"X-Demo-User-Id": "demo-rm"}
    )
    assert final_report.status_code == 200
    assert final_report.json()["report_status"] == "CONFIRMED"
    export = client.get(
        f"/api/v1/runs/{run_id}/report/export",
        headers={"X-Demo-User-Id": "demo-rm"},
    )
    assert export.status_code == 200
    assert export.headers["content-disposition"].endswith(f'credit-review-{run_id}.md"')
    assert "## 3. 规则结果" in export.text


def test_report_review_return_for_rerun_requires_reason(client: TestClient) -> None:
    case = create_case(client)
    document_ids = upload_required_documents(client, case["id"], valid=True)
    response = client.post(
        f"/api/v1/cases/{case['id']}/runs",
        headers=headers(key="run-report-return"),
        json={"document_version_ids": document_ids, "expected_case_version": 5},
    )
    assert response.status_code == 202
    run_id = response.json()["id"]
    from app.worker import Worker

    Worker(worker_id="return-worker", allow_memory_checkpoint=True).run_once()
    draft = client.get(
        f"/api/v1/runs/{run_id}/report", headers={"X-Demo-User-Id": "demo-reviewer"}
    ).json()
    missing_reason = client.post(
        f"/api/v1/runs/{run_id}/report-review",
        headers=headers(user="demo-reviewer", key="report-return-missing"),
        json={
            "expected_snapshot_version": draft["snapshot_version"],
            "action": "RETURN_FOR_RERUN",
        },
    )
    assert missing_reason.status_code == 400
    assert missing_reason.json()["code"] == "REASON_REQUIRED"
    returned_payload = {
        "expected_snapshot_version": draft["snapshot_version"],
        "action": "RETURN_FOR_RERUN",
        "reason": "需要补充最新财务报表。",
    }
    returned = client.post(
        f"/api/v1/runs/{run_id}/report-review",
        headers=headers(user="demo-reviewer", key="report-return-001"),
        json=returned_payload,
    )
    assert returned.status_code == 200
    assert returned.json()["status"] == "RETURNED"
    report = client.get(
        f"/api/v1/runs/{run_id}/report", headers={"X-Demo-User-Id": "demo-reviewer"}
    )
    assert report.status_code == 200
    assert report.json()["report_status"] == "RETURNED"
